"""Tests for the upload file converter (text passthrough, DOCX, PDF)."""

import io

import pytest

from brain_kb.convert import to_markdown


def test_text_passthrough():
    assert to_markdown("notes.md", b"# Title\nbody") == "# Title\nbody"


def test_docx_conversion_maps_headings():
    from docx import Document

    doc = Document()
    doc.add_heading("Quarterly Report", level=1)
    doc.add_paragraph("Revenue grew 20% this quarter.")
    doc.add_heading("Risks", level=2)
    doc.add_paragraph("Supply chain delays.")
    buf = io.BytesIO()
    doc.save(buf)

    md = to_markdown("report.docx", buf.getvalue())
    assert "# Quarterly Report" in md
    assert "## Risks" in md
    assert "Revenue grew 20% this quarter." in md


def test_docx_conversion_includes_tables():
    from docx import Document

    doc = Document()
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Team", "Budget"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Eng", "100k"
    buf = io.BytesIO()
    doc.save(buf)

    md = to_markdown("budget.docx", buf.getvalue())
    assert "| Team | Budget |" in md
    assert "| Eng | 100k |" in md


def test_pdf_conversion_extracts_text():
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 720, "Hello from a PDF document about billing.")
    c.save()

    md = to_markdown("doc.pdf", buf.getvalue())
    assert "Hello from a PDF document about billing." in md


def test_unsupported_binary_raises():
    with pytest.raises(ValueError):
        to_markdown("blob.bin", b"\xff\xfe\x00\x01\x02")
